from __future__ import annotations

import re
from dataclasses import dataclass, field
from html import escape, unescape
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook


OBJECT_PREFIXES = {
    "NEED": "Need",
    "REQ": "Requirement",
    "FUN": "Function",
    "PART": "Part",
    "CON": "Concept",
    "VER": "Verification",
    "DEC": "Decision",
    "RISK": "Risk",
    "ISSUE": "Issue",
    "TASK": "Task",
    "MS": "Milestone",
}

OBJECT_ID_PATTERN = re.compile(r"^(?P<prefix>[A-Z]+)-[0-9]{3}$")


@dataclass(frozen=True, slots=True)
class SourceRef:
    """Logical and physical provenance for parsed Markdown content.

    ``file``, ``table_index``, ``row_id`` and ``column`` form the reusable
    logical locator. ``row_index`` and ``line`` are useful diagnostics captured
    at load time, but writers must re-resolve the row by its stable ID.
    """

    file: str
    table_index: int
    row_index: int | None = None
    row_id: str | None = None
    column: str | None = None
    line: int | None = None

    def for_column(self, column: str) -> SourceRef:
        return SourceRef(
            file=self.file,
            table_index=self.table_index,
            row_index=self.row_index,
            row_id=self.row_id,
            column=column,
            line=self.line,
        )

    def to_dict(self) -> dict[str, str | int | None]:
        return {
            "file": self.file,
            "table_index": self.table_index,
            "row_index": self.row_index,
            "row_id": self.row_id,
            "column": self.column,
            "line": self.line,
        }


@dataclass(frozen=True, slots=True)
class ParsedMarkdownTable:
    """One lightweight Markdown table plus provenance for each parsed row."""

    headers: tuple[str, ...]
    rows: list[dict[str, str]]
    source_ref: SourceRef
    row_sources: tuple[SourceRef, ...]


@dataclass(slots=True)
class ModelObject:
    id: str
    type: str
    attributes: dict[str, str] = field(default_factory=dict)
    source_file: str = ""
    source_ref: SourceRef | None = None
    attribute_sources: dict[str, SourceRef] = field(default_factory=dict)


@dataclass(slots=True)
class Relation:
    source: str
    relation: str
    target: str
    source_file: str = ""
    source_ref: SourceRef | None = None


@dataclass(slots=True)
class Model:
    objects: dict[str, ModelObject] = field(default_factory=dict)
    relations: list[Relation] = field(default_factory=list)
    tables: dict[str, list[dict[str, str]]] = field(default_factory=dict)
    table_sources: dict[str, SourceRef] = field(default_factory=dict)
    table_row_sources: dict[str, tuple[SourceRef, ...]] = field(default_factory=dict)
    duplicate_ids: list[str] = field(default_factory=list)


def _split_markdown_row(line: str) -> list[str]:
    row = line.strip()
    if row.startswith("|"):
        row = row[1:]
    if row.endswith("|"):
        row = row[:-1]
    return [unescape(cell.strip()) for cell in row.split("|")]


def _is_separator_row(cells: list[str]) -> bool:
    return bool(cells) and all(cell.replace(":", "").replace("-", "").strip() == "" for cell in cells)


def parse_markdown_tables_with_provenance(
    path: Path, *, source_file: str | None = None
) -> list[ParsedMarkdownTable]:
    """Parse simple Markdown tables without constructing a full Markdown AST."""

    lines = path.read_text(encoding="utf-8-sig").splitlines()
    file_name = source_file if source_file is not None else path.as_posix()
    tables: list[ParsedMarkdownTable] = []
    i = 0
    while i + 1 < len(lines):
        if "|" not in lines[i] or "|" not in lines[i + 1]:
            i += 1
            continue
        headers = _split_markdown_row(lines[i])
        separator = _split_markdown_row(lines[i + 1])
        if len(headers) < 2 or len(headers) != len(separator) or not _is_separator_row(separator):
            i += 1
            continue
        header_line = i + 1
        rows: list[dict[str, str]] = []
        row_sources: list[SourceRef] = []
        i += 2
        while i < len(lines) and "|" in lines[i] and lines[i].strip():
            cells = _split_markdown_row(lines[i])
            if len(cells) != len(headers):
                break
            row = dict(zip(headers, cells))
            rows.append(row)
            row_sources.append(
                SourceRef(
                    file=file_name,
                    table_index=len(tables) + 1,
                    row_index=len(rows),
                    row_id=row.get("ID", "").strip() or None,
                    line=i + 1,
                )
            )
            i += 1
        if rows:
            tables.append(
                ParsedMarkdownTable(
                    headers=tuple(headers),
                    rows=rows,
                    source_ref=SourceRef(
                        file=file_name,
                        table_index=len(tables) + 1,
                        line=header_line,
                    ),
                    row_sources=tuple(row_sources),
                )
            )
    return tables


def parse_markdown_tables(path: Path) -> list[list[dict[str, str]]]:
    """Backward-compatible table-only parser."""

    return [table.rows for table in parse_markdown_tables_with_provenance(path)]


def load_model(project_dir: str | Path) -> Model:
    project_path = Path(project_dir)
    model = Model()
    for md_path in sorted(project_path.rglob("*.md")):
        relative_path = md_path.relative_to(project_path).as_posix()
        parsed_tables = parse_markdown_tables_with_provenance(
            md_path, source_file=relative_path
        )
        for table in parsed_tables:
            rows = table.rows
            key = f"{relative_path}:{table.source_ref.table_index}"
            model.tables[key] = rows
            model.table_sources[key] = table.source_ref
            model.table_row_sources[key] = table.row_sources
            headers = set(rows[0]) if rows else set()
            if {"Source", "Relation", "Target"}.issubset(headers):
                for row, row_source in zip(rows, table.row_sources):
                    model.relations.append(
                        Relation(
                            source=row.get("Source", "").strip(),
                            relation=row.get("Relation", "").strip(),
                            target=row.get("Target", "").strip(),
                            source_file=relative_path,
                            source_ref=row_source,
                        )
                    )
                continue
            if "ID" not in headers:
                continue
            for row, row_source in zip(rows, table.row_sources):
                object_id = row.get("ID", "").strip()
                if not object_id:
                    continue
                prefix = object_id.split("-", 1)[0]
                object_type = OBJECT_PREFIXES.get(prefix, prefix or "Unknown")
                obj = ModelObject(
                    id=object_id,
                    type=object_type,
                    attributes={k: v for k, v in row.items() if k != "ID"},
                    source_file=relative_path,
                    source_ref=row_source,
                    attribute_sources={
                        column: row_source.for_column(column)
                        for column in row
                        if column != "ID"
                    },
                )
                if object_id in model.objects:
                    model.duplicate_ids.append(object_id)
                else:
                    model.objects[object_id] = obj
    return model


def validate_model(model: Model) -> list[tuple[str, str]]:
    findings: list[tuple[str, str]] = []
    for object_id in model.duplicate_ids:
        findings.append(("ERROR", f"Duplicate ID: {object_id}"))

    for obj in model.objects.values():
        match = OBJECT_ID_PATTERN.fullmatch(obj.id)
        if match is None:
            findings.append(("ERROR", f"Invalid ID format: {obj.id}; expected PREFIX-NNN"))
        elif match.group("prefix") not in OBJECT_PREFIXES:
            findings.append(("ERROR", f"Unknown ID prefix: {match.group('prefix')} in {obj.id}"))

    for rel in model.relations:
        if rel.source not in model.objects:
            findings.append(("ERROR", f"Unknown relation source {rel.source} in {rel.source_file}"))
        if rel.target not in model.objects:
            findings.append(("ERROR", f"Unknown relation target {rel.target} in {rel.source_file}"))
        if not rel.relation:
            findings.append(("ERROR", f"Empty relation type for {rel.source} -> {rel.target}"))

    outgoing: dict[str, list[Relation]] = {}
    incoming: dict[str, list[Relation]] = {}
    for rel in model.relations:
        outgoing.setdefault(rel.source, []).append(rel)
        incoming.setdefault(rel.target, []).append(rel)

    for obj in model.objects.values():
        if obj.type == "Requirement":
            has_need_derivation = any(
                rel.relation == "derives"
                and (source := model.objects.get(rel.source)) is not None
                and source.type == "Need"
                for rel in incoming.get(obj.id, [])
            )
            if not has_need_derivation:
                findings.append(("WARNING", f"Requirement {obj.id} has no incoming derives relation from a Need"))

            requirement_relations = {rel.relation for rel in outgoing.get(obj.id, [])}
            if "satisfied_by" not in requirement_relations:
                findings.append(("WARNING", f"Requirement {obj.id} has no outgoing satisfied_by relation"))
            if "verified_by" not in requirement_relations:
                findings.append(("WARNING", f"Requirement {obj.id} has no outgoing verified_by relation"))

        if obj.type == "Function" and not any(
            rel.relation == "realized_by" for rel in outgoing.get(obj.id, [])
        ):
            findings.append(("WARNING", f"Function {obj.id} has no outgoing realized_by relation"))

    return findings


def traceability_rows(model: Model) -> list[dict[str, str]]:
    return [
        {"Source": rel.source, "Relation": rel.relation, "Target": rel.target}
        for rel in model.relations
    ]


def mermaid_graph(model: Model) -> str:
    lines = ["flowchart LR"]
    for obj in model.objects.values():
        label = obj.attributes.get("Name") or obj.attributes.get("Title") or obj.id
        safe_label = label.replace('"', "'")
        lines.append(f'    {obj.id.replace("-", "_")}["{obj.id}<br/>{safe_label}"]')
    for rel in model.relations:
        if rel.source in model.objects and rel.target in model.objects:
            lines.append(
                f'    {rel.source.replace("-", "_")} -->|{rel.relation}| {rel.target.replace("-", "_")}'
            )
    return "\n".join(lines)


def export_mermaid(model: Model, output: str | Path) -> None:
    output_path = Path(output)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(f"# Traceability view\n\n```mermaid\n{mermaid_graph(model)}\n```\n", encoding="utf-8")


def export_xlsx(model: Model, output: str | Path) -> None:
    output_path = Path(output)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)

    by_type: dict[str, list[ModelObject]] = {}
    for obj in model.objects.values():
        by_type.setdefault(obj.type, []).append(obj)

    for object_type, objects in sorted(by_type.items()):
        ws = wb.create_sheet(title=object_type[:31])
        attribute_names: list[str] = sorted({key for obj in objects for key in obj.attributes})
        ws.append(["ID", "Type", *attribute_names, "Source File"])
        for obj in objects:
            ws.append([obj.id, obj.type, *[obj.attributes.get(name, "") for name in attribute_names], obj.source_file])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    ws = wb.create_sheet(title="Relations")
    ws.append(["Source", "Relation", "Target", "Source File"])
    for rel in model.relations:
        ws.append([rel.source, rel.relation, rel.target, rel.source_file])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    findings = validate_model(model)
    ws = wb.create_sheet(title="Validation")
    ws.append(["Severity", "Finding"])
    for severity, message in findings:
        ws.append([severity, message])

    wb.save(output_path)


def _markdown_cell(value: object) -> str:
    if value is None:
        return ""
    return (
        str(value)
        .replace("&", "&amp;")
        .replace("\r", "&#13;")
        .replace("\n", "&#10;")
        .replace("|", "&#124;")
        .strip()
    )


def import_xlsx_to_markdown(input_file: str | Path, output_dir: str | Path) -> list[Path]:
    """Convert an MBSE-lite XLSX export to normalized Markdown files for review.

    This is deliberately a safe import: it never edits an existing Markdown project in place.
    The generated directory should be reviewed before replacing source-of-truth files.
    """
    input_path = Path(input_file)
    out = Path(output_dir)
    out.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(input_path, data_only=False)
    written: list[Path] = []

    for ws in wb.worksheets:
        if ws.title == "Validation" or ws.max_row < 1:
            continue
        headers = [_markdown_cell(cell.value) for cell in ws[1]]
        if not any(headers):
            continue

        drop = {"Type", "Source File"}
        keep_indices = [i for i, header in enumerate(headers) if header and header not in drop]
        kept_headers = [headers[i] for i in keep_indices]
        if not kept_headers:
            continue

        rows: list[list[str]] = []
        for excel_row in ws.iter_rows(min_row=2, values_only=True):
            row = [_markdown_cell(excel_row[i]) for i in keep_indices]
            if any(row):
                rows.append(row)

        file_name = "relations.md" if ws.title == "Relations" else f"{ws.title.lower().replace(' ', '_')}.md"
        path = out / file_name
        lines = [f"# {ws.title}", "", "| " + " | ".join(kept_headers) + " |", "|" + "|".join("---" for _ in kept_headers) + "|"]
        lines.extend("| " + " | ".join(row) + " |" for row in rows)
        lines.append("")
        path.write_text("\n".join(lines), encoding="utf-8")
        written.append(path)

    return written


def _html_table(headers: Iterable[str], rows: Iterable[Iterable[str]]) -> str:
    header_html = "".join(f"<th>{escape(str(h))}</th>" for h in headers)
    row_html = "".join(
        "<tr>" + "".join(f"<td>{escape(str(cell))}</td>" for cell in row) + "</tr>"
        for row in rows
    )
    return f"<table><thead><tr>{header_html}</tr></thead><tbody>{row_html}</tbody></table>"


def export_html(model: Model, output: str | Path) -> None:
    output_path = Path(output)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    findings = validate_model(model)
    counts: dict[str, int] = {}
    for obj in model.objects.values():
        counts[obj.type] = counts.get(obj.type, 0) + 1

    count_rows = [[key, str(value)] for key, value in sorted(counts.items())]
    validation_rows = [[severity, message] for severity, message in findings] or [["OK", "No findings"]]
    relation_rows = [[r.source, r.relation, r.target] for r in model.relations]

    html = f"""<!doctype html>
<html lang=\"en\">
<head>
<meta charset=\"utf-8\">
<title>MBSE Lite Project Overview</title>
<style>
body {{ font-family: system-ui, sans-serif; max-width: 1200px; margin: 2rem auto; padding: 0 1rem; }}
table {{ border-collapse: collapse; width: 100%; margin-bottom: 2rem; }}
th, td {{ border: 1px solid #ccc; padding: .45rem .6rem; text-align: left; }}
th {{ background: #f3f3f3; }}
pre {{ overflow-x: auto; background: #f6f6f6; padding: 1rem; }}
</style>
<script type=\"module\">
import mermaid from 'https://cdn.jsdelivr.net/npm/mermaid@11/dist/mermaid.esm.min.mjs';
mermaid.initialize({{ startOnLoad: true }});
</script>
</head>
<body>
<h1>MBSE Lite Project Overview</h1>
<h2>Model content</h2>
{_html_table(["Object type", "Count"], count_rows)}
<h2>Validation</h2>
{_html_table(["Severity", "Finding"], validation_rows)}
<h2>Traceability graph</h2>
<pre class=\"mermaid\">{escape(mermaid_graph(model))}</pre>
<h2>Relations</h2>
{_html_table(["Source", "Relation", "Target"], relation_rows)}
</body>
</html>
"""
    output_path.write_text(html, encoding="utf-8")
