from __future__ import annotations

from dataclasses import dataclass, field
from html import escape
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook


OBJECT_PREFIXES = {
    "NEED": "Need",
    "REQ": "Requirement",
    "FUN": "Function",
    "PART": "Part",
    "CON": "Concept",
    "VER": "Verification",
    "DEC": "Decision",
    "RISK": "Risk",
    "ISSUE": "Issue",
    "TASK": "Task",
    "MS": "Milestone",
}


@dataclass(slots=True)
class ModelObject:
    id: str
    type: str
    attributes: dict[str, str] = field(default_factory=dict)
    source_file: str = ""


@dataclass(slots=True)
class Relation:
    source: str
    relation: str
    target: str
    source_file: str = ""


@dataclass(slots=True)
class Model:
    objects: dict[str, ModelObject] = field(default_factory=dict)
    relations: list[Relation] = field(default_factory=list)
    tables: dict[str, list[dict[str, str]]] = field(default_factory=dict)
    duplicate_ids: list[str] = field(default_factory=list)


def _split_markdown_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def _is_separator_row(cells: list[str]) -> bool:
    return bool(cells) and all(cell.replace(":", "").replace("-", "").strip() == "" for cell in cells)


def parse_markdown_tables(path: Path) -> list[list[dict[str, str]]]:
    lines = path.read_text(encoding="utf-8").splitlines()
    tables: list[list[dict[str, str]]] = []
    i = 0
    while i + 1 < len(lines):
        if "|" not in lines[i] or "|" not in lines[i + 1]:
            i += 1
            continue
        headers = _split_markdown_row(lines[i])
        separator = _split_markdown_row(lines[i + 1])
        if len(headers) < 2 or len(headers) != len(separator) or not _is_separator_row(separator):
            i += 1
            continue
        rows: list[dict[str, str]] = []
        i += 2
        while i < len(lines) and "|" in lines[i] and lines[i].strip():
            cells = _split_markdown_row(lines[i])
            if len(cells) != len(headers):
                break
            rows.append(dict(zip(headers, cells)))
            i += 1
        if rows:
            tables.append(rows)
    return tables


def load_model(project_dir: str | Path) -> Model:
    project_path = Path(project_dir)
    model = Model()
    for md_path in sorted(project_path.glob("*.md")):
        parsed_tables = parse_markdown_tables(md_path)
        for index, rows in enumerate(parsed_tables, start=1):
            key = f"{md_path.stem}:{index}"
            model.tables[key] = rows
            headers = set(rows[0]) if rows else set()
            if {"Source", "Relation", "Target"}.issubset(headers):
                for row in rows:
                    model.relations.append(
                        Relation(
                            source=row.get("Source", "").strip(),
                            relation=row.get("Relation", "").strip(),
                            target=row.get("Target", "").strip(),
                            source_file=md_path.name,
                        )
                    )
                continue
            if "ID" not in headers:
                continue
            for row in rows:
                object_id = row.get("ID", "").strip()
                if not object_id:
                    continue
                prefix = object_id.split("-", 1)[0]
                object_type = OBJECT_PREFIXES.get(prefix, prefix or "Unknown")
                obj = ModelObject(
                    id=object_id,
                    type=object_type,
                    attributes={k: v for k, v in row.items() if k != "ID"},
                    source_file=md_path.name,
                )
                if object_id in model.objects:
                    model.duplicate_ids.append(object_id)
                else:
                    model.objects[object_id] = obj
    return model


def validate_model(model: Model) -> list[tuple[str, str]]:
    findings: list[tuple[str, str]] = []
    for object_id in model.duplicate_ids:
        findings.append(("ERROR", f"Duplicate ID: {object_id}"))

    for rel in model.relations:
        if rel.source not in model.objects:
            findings.append(("ERROR", f"Unknown relation source {rel.source} in {rel.source_file}"))
        if rel.target not in model.objects:
            findings.append(("ERROR", f"Unknown relation target {rel.target} in {rel.source_file}"))
        if not rel.relation:
            findings.append(("ERROR", f"Empty relation type for {rel.source} -> {rel.target}"))

    outgoing: dict[str, list[Relation]] = {}
    incoming: dict[str, list[Relation]] = {}
    for rel in model.relations:
        outgoing.setdefault(rel.source, []).append(rel)
        incoming.setdefault(rel.target, []).append(rel)

    for obj in model.objects.values():
        if obj.type == "Requirement":
            if not incoming.get(obj.id):
                findings.append(("WARNING", f"Requirement {obj.id} has no incoming traceability relation"))
            if not outgoing.get(obj.id):
                findings.append(("WARNING", f"Requirement {obj.id} has no outgoing traceability relation"))
            if not any(rel.relation in {"verified_by", "verify", "verifiedBy"} for rel in outgoing.get(obj.id, [])):
                findings.append(("WARNING", f"Requirement {obj.id} has no verification relation"))
        if obj.type == "Function" and not outgoing.get(obj.id):
            findings.append(("WARNING", f"Function {obj.id} is not realized by another model object"))

    return findings


def traceability_rows(model: Model) -> list[dict[str, str]]:
    return [
        {"Source": rel.source, "Relation": rel.relation, "Target": rel.target}
        for rel in model.relations
    ]


def mermaid_graph(model: Model) -> str:
    lines = ["flowchart LR"]
    for obj in model.objects.values():
        label = obj.attributes.get("Name") or obj.attributes.get("Title") or obj.id
        safe_label = label.replace('"', "'")
        lines.append(f'    {obj.id.replace("-", "_")}["{obj.id}<br/>{safe_label}"]')
    for rel in model.relations:
        if rel.source in model.objects and rel.target in model.objects:
            lines.append(
                f'    {rel.source.replace("-", "_")} -->|{rel.relation}| {rel.target.replace("-", "_")}'
            )
    return "\n".join(lines)


def export_mermaid(model: Model, output: str | Path) -> None:
    output_path = Path(output)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(f"# Traceability view\n\n```mermaid\n{mermaid_graph(model)}\n```\n", encoding="utf-8")


def export_xlsx(model: Model, output: str | Path) -> None:
    output_path = Path(output)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)

    by_type: dict[str, list[ModelObject]] = {}
    for obj in model.objects.values():
        by_type.setdefault(obj.type, []).append(obj)

    for object_type, objects in sorted(by_type.items()):
        ws = wb.create_sheet(title=object_type[:31])
        attribute_names: list[str] = sorted({key for obj in objects for key in obj.attributes})
        ws.append(["ID", "Type", *attribute_names, "Source File"])
        for obj in objects:
            ws.append([obj.id, obj.type, *[obj.attributes.get(name, "") for name in attribute_names], obj.source_file])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    ws = wb.create_sheet(title="Relations")
    ws.append(["Source", "Relation", "Target", "Source File"])
    for rel in model.relations:
        ws.append([rel.source, rel.relation, rel.target, rel.source_file])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    findings = validate_model(model)
    ws = wb.create_sheet(title="Validation")
    ws.append(["Severity", "Finding"])
    for severity, message in findings:
        ws.append([severity, message])

    wb.save(output_path)


def _html_table(headers: Iterable[str], rows: Iterable[Iterable[str]]) -> str:
    header_html = "".join(f"<th>{escape(str(h))}</th>" for h in headers)
    row_html = "".join(
        "<tr>" + "".join(f"<td>{escape(str(cell))}</td>" for cell in row) + "</tr>"
        for row in rows
    )
    return f"<table><thead><tr>{header_html}</tr></thead><tbody>{row_html}</tbody></table>"


def export_html(model: Model, output: str | Path) -> None:
    output_path = Path(output)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    findings = validate_model(model)
    counts: dict[str, int] = {}
    for obj in model.objects.values():
        counts[obj.type] = counts.get(obj.type, 0) + 1

    count_rows = [[key, str(value)] for key, value in sorted(counts.items())]
    validation_rows = [[severity, message] for severity, message in findings] or [["OK", "No findings"]]
    relation_rows = [[r.source, r.relation, r.target] for r in model.relations]

    html = f"""<!doctype html>
<html lang=\"en\">
<head>
<meta charset=\"utf-8\">
<title>MBSE Lite Project Overview</title>
<style>
body {{ font-family: system-ui, sans-serif; max-width: 1200px; margin: 2rem auto; padding: 0 1rem; }}
table {{ border-collapse: collapse; width: 100%; margin-bottom: 2rem; }}
th, td {{ border: 1px solid #ccc; padding: .45rem .6rem; text-align: left; }}
th {{ background: #f3f3f3; }}
pre {{ overflow-x: auto; background: #f6f6f6; padding: 1rem; }}
</style>
</head>
<body>
<h1>MBSE Lite Project Overview</h1>
<h2>Model content</h2>
{_html_table(["Object type", "Count"], count_rows)}
<h2>Validation</h2>
{_html_table(["Severity", "Finding"], validation_rows)}
<h2>Relations</h2>
{_html_table(["Source", "Relation", "Target"], relation_rows)}
<h2>Mermaid source</h2>
<p>Paste this block into a Mermaid-capable Markdown viewer to render the graph.</p>
<pre>{escape(mermaid_graph(model))}</pre>
</body>
</html>
"""
    output_path.write_text(html, encoding="utf-8")
